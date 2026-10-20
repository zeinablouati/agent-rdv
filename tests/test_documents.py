"""Documents : le RAG lit aussi les PDF et Word, pas seulement les .md."""

from agent_rdv.knowledge.base import _decouper_par_taille, charger_passages


def test_decoupe_par_taille_avec_chevauchement():
    texte = "mot " * 500  # ~2000 caractères
    morceaux = _decouper_par_taille(texte, taille=800, chevauchement=100)
    assert len(morceaux) >= 2
    assert all(len(m) <= 800 for m in morceaux)


def test_charger_md_et_txt(tmp_path):
    (tmp_path / "info.md").write_text(
        "# Titre général\n\n## Tarifs\n30 euros.\n\n## Horaires\nLundi au vendredi.", encoding="utf-8"
    )
    (tmp_path / "note.txt").write_text("Paragraphe un.\n\nParagraphe deux.", encoding="utf-8")
    textes = [t for _, t in charger_passages(tmp_path)]
    assert any("30 euros" in t for t in textes)
    assert any("Paragraphe deux" in t for t in textes)
    assert not any(t.startswith("# Titre") for t in textes)  # le grand titre est ignoré


def test_charger_docx(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("Le cabinet accepte la carte bancaire.")
    document.add_paragraph("Stationnement payant à proximité.")
    document.save(str(tmp_path / "pratique.docx"))

    textes = " ".join(t for _, t in charger_passages(tmp_path))
    assert "carte bancaire" in textes


def test_ignore_les_extensions_non_prises_en_charge(tmp_path):
    (tmp_path / "image.png").write_bytes(b"\x89PNG...")
    (tmp_path / "info.md").write_text("## Accès\nMétro ligne 12.", encoding="utf-8")
    noms = [pid for pid, _ in charger_passages(tmp_path)]
    assert any(n.startswith("info.md") for n in noms)
    assert not any("image" in n for n in noms)